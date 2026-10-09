from pathlib import Path

from openpyxl import load_workbook

from modelo import FilaEntrada

ENCABEZADOS = ("IdLista", "Factura", "NIT", "Proveedor", "Valor", "CentroCosto", "FechaFactura")


def _normalizar(texto: object) -> str:
    return str(texto).strip().lower() if texto is not None else ""


def _vacio(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def leer_por_radicar(ruta: Path) -> list[FilaEntrada]:
    """Lee la primera hoja del libro ubicando las columnas por su encabezado."""
    libro = load_workbook(ruta, data_only=True, read_only=True)
    try:
        hoja = libro.worksheets[0]
        filas = hoja.iter_rows(values_only=True)

        columnas = None
        for fila in filas:
            nombres = [_normalizar(c) for c in fila]
            if all(e.lower() in nombres for e in ENCABEZADOS):
                columnas = {e: nombres.index(e.lower()) for e in ENCABEZADOS}
                break
        if columnas is None:
            raise SystemExit(f"{ruta.name}: la primera hoja no tiene los encabezados {', '.join(ENCABEZADOS)}")

        entradas = []
        for fila in filas:
            def celda(nombre: str) -> object:
                i = columnas[nombre]
                return fila[i] if i < len(fila) else None

            if _vacio(celda("IdLista")):
                continue
            entradas.append(
                FilaEntrada(
                    id_lista=celda("IdLista"),
                    factura=celda("Factura"),
                    nit=celda("NIT"),
                    proveedor=celda("Proveedor"),
                    valor=celda("Valor"),
                    centro_costo=celda("CentroCosto"),
                    fecha_factura=celda("FechaFactura"),
                )
            )
        return entradas
    finally:
        libro.close()
